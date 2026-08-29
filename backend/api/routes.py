from fastapi import APIRouter, UploadFile, File, HTTPException, status
from fastapi.responses import Response, StreamingResponse
import io
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime

from models.contact import (
    ContactModel,
    ContactScanResult,
    ConfigStatusResponse,
    SaveContactResponse,
    MergeContactsRequest,
    RenameLedgerRequest
)
from utils.validation import validate_image_file
from services.image_processing import process_image
from services.ocr import run_ocr
from services.extractor import extract_contact_info
from services.google_sheets import (
    is_google_sheets_configured,
    get_contact_count,
    append_contact_to_sheet,
    archive_active_worksheet_in_gs,
    activate_archived_worksheet_in_gs,
    delete_archived_worksheet_in_gs,
    rename_archived_worksheet_in_gs,
    get_all_contacts_from_sheet,
    update_contact_in_sheet,
    delete_contact_from_sheet,
    merge_contacts_in_sheet,
    get_all_ledgers_from_sheets,
    filename_to_tab_name,
    HEADERS
)
from services.excel_storage import generate_vcard_string

import logging
logger = logging.getLogger("cardsnap.routes")

router = APIRouter()

@router.get("/contacts")
def list_contacts():
    contacts = get_all_contacts_from_sheet()
    return {"contacts": contacts, "total": len(contacts)}

@router.put("/contacts/{row_id}")
def update_contact(row_id: int, contact: ContactModel):
    success, msg = update_contact_in_sheet(row_id, contact)
    if not success:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=msg)
    return {"success": True, "message": msg}

@router.delete("/contacts/{row_id}")
def delete_contact(row_id: int):
    success, msg = delete_contact_from_sheet(row_id)
    if not success:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=msg)
    return {"success": True, "message": msg}

@router.post("/contacts/reset")
def reset_contacts_sheet():
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    archive_filename = f"CardSnap_Contacts_{timestamp}.xlsx"
    
    success, msg = archive_active_worksheet_in_gs(archive_filename)
    if not success:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=msg)
        
    return {"success": True, "message": msg}

@router.get("/download-vcard")
def download_vcard():
    contacts = get_all_contacts_from_sheet()
    vcard_str = generate_vcard_string(contacts)
    return Response(
        content=vcard_str,
        media_type="text/vcard",
        headers={"Content-Disposition": "attachment; filename=CardSnap_Contacts.vcf"}
    )

@router.get("/health")
def health_check():
    return {"status": "ok"}

@router.get("/config/status", response_model=ConfigStatusResponse)
def get_config_status():
    try:
        is_configured, sheet_id, message = is_google_sheets_configured()
        count = get_contact_count() if is_configured else None
        return ConfigStatusResponse(
            google_sheets_configured=is_configured,
            spreadsheet_id=sheet_id if is_configured else None,
            contact_count=count,
            local_excel_count=count,
            message=message or ""
        )
    except Exception as e:
        logger.error(f"Error in get_config_status: {e}", exc_info=True)
        return ConfigStatusResponse(
            google_sheets_configured=False,
            spreadsheet_id=None,
            contact_count=None,
            local_excel_count=None,
            message=f"Configuration status check error: {str(e)}"
        )

@router.get("/download-excel")
def download_excel():
    contacts = get_all_contacts_from_sheet()
    
    # Create workbook in memory
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"
    
    ws.append(HEADERS)
    
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    
    for col_num, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        ws.column_dimensions[get_column_letter(col_num)].width = max(len(header) + 6, 18)
        
    for c in contacts:
        ws.append([
            c.get("name") or "",
            c.get("job_title") or "",
            c.get("company") or "",
            c.get("phone") or "",
            c.get("email") or "",
            c.get("website") or "",
            c.get("linkedin") or "",
            c.get("address") or "",
            c.get("notes") or "",
            c.get("date_added") or ""
        ])
        
    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=CardSnap_Contacts.xlsx"}
    )

@router.get("/ledgers")
def list_ledgers():
    return {"ledgers": get_all_ledgers_from_sheets()}

@router.post("/ledgers/activate/{filename}")
def activate_ledger(filename: str):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_filename = f"CardSnap_Contacts_{timestamp}.xlsx"
    
    success, msg = activate_archived_worksheet_in_gs(filename, backup_filename)
    if not success:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=msg)
            
    return {"success": True, "message": msg}

@router.delete("/ledgers/{filename}")
def delete_ledger(filename: str):
    success, msg = delete_archived_worksheet_in_gs(filename)
    if not success:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=msg)
            
    return {"success": True, "message": msg}

@router.get("/download-excel/{filename}")
def download_specific_excel(filename: str):
    tab_name = filename_to_tab_name(filename)
    
    is_conf, sheet_id, msg = is_google_sheets_configured()
    if not is_conf or not sheet_id:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Google Sheet not configured")
        
    try:
        client = get_gspread_client()
        sh = client.open_by_key(sheet_id)
        wks = sh.worksheet(tab_name)
        all_vals = wks.get_all_values()
    except Exception as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Ledger tab '{tab_name}' not found: {str(e)}")
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"
    
    for row in all_vals:
        ws.append(row)
        
    if all_vals:
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        for col_num in range(1, len(all_vals[0]) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            ws.column_dimensions[get_column_letter(col_num)].width = 18
              
    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.post("/scan", response_model=ContactScanResult)
async def scan_business_card(file: UploadFile = File(...)):
    try:
        contents = await file.read()
        validate_image_file(file, contents)
        
        # Image processing
        processed_img, original_img = process_image(contents)
        
        # OCR execution
        ocr_results = run_ocr(processed_img)
        if not ocr_results:
            ocr_results = run_ocr(original_img)
            
        if not ocr_results:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="We couldn't read this card clearly. Try taking a sharper, better-lit photo."
            )
            
        extracted_result = extract_contact_info(ocr_results)
        extracted_result.card_image_filename = ""
        return extracted_result

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error scanning business card: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to process business card: {str(e)}"
        )

@router.post("/save-contact", response_model=SaveContactResponse)
def save_contact(contact: ContactModel):
    gs_success, gs_msg, gs_row = append_contact_to_sheet(contact)
    if not gs_success:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=gs_msg
        )
    return SaveContactResponse(
        success=True,
        message="Saved to Google Sheet successfully.",
        row_added=gs_row
    )

@router.get("/card-image/{filename}")
def get_card_image(filename: str):
    raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Image storage disabled in cloud environment.")

@router.get("/contacts/duplicates")
def get_duplicate_contacts():
    import re
    import difflib
    contacts = get_all_contacts_from_sheet()
    if not contacts:
        return {"duplicates": []}
        
    n = len(contacts)
    parent = list(range(n))
    
    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
        
    def union(i, j):
        root_i = find(i)
        root_j = find(j)
        if root_i != root_j:
            parent[root_i] = root_j
            
    for i in range(n):
        c1 = contacts[i]
        email1 = c1.get("email", "").strip().lower()
        phone1 = re.sub(r'\D', '', c1.get("phone", ""))
        name1 = c1.get("name", "").strip().lower()
        
        for j in range(i + 1, n):
            c2 = contacts[j]
            email2 = c2.get("email", "").strip().lower()
            phone2 = re.sub(r'\D', '', c2.get("phone", ""))
            name2 = c2.get("name", "").strip().lower()
            
            is_match = False
            
            if email1 and email2 and email1 == email2:
                is_match = True
            elif phone1 and phone2 and len(phone1) >= 7 and phone1 == phone2:
                is_match = True
            elif name1 and name2 and len(name1) > 2 and len(name2) > 2:
                if name1 == name2:
                    is_match = True
                else:
                    ratio = difflib.SequenceMatcher(None, name1, name2).ratio()
                    if ratio >= 0.85:
                        is_match = True
            
            if is_match:
                union(i, j)
                
    clusters = {}
    for i in range(n):
        root = find(i)
        clusters.setdefault(root, []).append(contacts[i])
        
    duplicate_groups = []
    for root, group in clusters.items():
        if len(group) > 1:
            email1 = group[0].get("email", "").strip().lower()
            phone1 = re.sub(r'\D', '', group[0].get("phone", ""))
            name1 = group[0].get("name", "").strip().lower()
            
            m_type = "Name"
            m_val = group[0].get("name")
            
            for k in range(1, len(group)):
                email2 = group[k].get("email", "").strip().lower()
                phone2 = re.sub(r'\D', '', group[k].get("phone", ""))
                name2 = group[k].get("name", "").strip().lower()
                
                if email1 and email2 and email1 == email2:
                    m_type = "Email"
                    m_val = email1
                    break
                elif phone1 and phone2 and len(phone1) >= 7 and phone1 == phone2:
                    m_type = "Phone"
                    m_val = group[0].get("phone")
                    break
                elif name1 and name2 and len(name1) > 2 and len(name2) > 2:
                    if name1 != name2:
                        m_type = "Fuzzy Name"
                        m_val = f"{group[0].get('name')} ~ {group[k].get('name')}"
                        break
                    
            duplicate_groups.append({
                "type": m_type,
                "value": m_val,
                "contacts": group
            })
            
    return {"duplicates": duplicate_groups}

@router.post("/contacts/merge")
def merge_contacts(req: MergeContactsRequest):
    success, msg = merge_contacts_in_sheet(
        target_row_id=req.target_row_id,
        duplicate_row_ids=req.duplicate_row_ids,
        merged_contact=req.merged_contact
    )
    if not success:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=msg)
    return {"success": True, "message": msg}

@router.post("/ledgers/rename")
def rename_ledger(req: RenameLedgerRequest):
    clean_label = "".join(c for c in req.new_label if c.isalnum() or c in ("_", "-")).strip()
    if not clean_label:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid label name.")
    new_filename = f"CardSnap_Contacts_{clean_label}.xlsx"
    
    success, msg = rename_archived_worksheet_in_gs(req.filename, new_filename)
    if not success:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=msg)
    return {"success": True, "message": msg}


